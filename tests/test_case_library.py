from io import BytesIO
import json

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app.main import create_app
from backend.app.services.case_library_service import CaseLibraryService, excel_bytes


def test_library_matches_visit_bindings_and_canonical_truth(cases):
    library = CaseLibraryService(cases)
    assert library.index()["case_count"] == len(cases.cases)
    for cid in cases.case_ids():
        detail = library.detail(cid)
        assert {a["id"] for a in detail["actions"]} == set(library.catalog.bindings[cid])
        assert {e["evidence_id"] for e in detail["model_evidence"]} == set(cases.catalog_for_case(cid))
        for evidence in detail["model_evidence"]:
            assert evidence["source_response"] == cases.get(cid).truth_map()[evidence["evidence_id"]].model_dump(mode="json")
        assert detail["initial_information"] == library.catalog.inspection(cid, "en")["visit"]
        assert "oracle" not in detail["initial_information"]


def test_library_exports_are_protected_and_scoped(cases):
    client = TestClient(create_app(cases))
    cid = cases.case_ids()[0]
    for path in ["/research/case-library", f"/research/case-library/{cid}",
                 "/research/case-library/export"]:
        assert client.get(path).status_code == 403
    headers = {"X-Research-Key": "local-research-only"}
    response = client.get("/research/case-library/export", params={"case_id": cid}, headers=headers)
    assert response.status_code == 200
    assert response.json()["case_count"] == 1
    assert response.json()["cases"][0]["case_id"] == cid
    assert "attachment" in response.headers["content-disposition"]
    assert client.get("/research/case-library/export?format=csv", headers=headers).status_code == 400
    assert client.get("/research/case-library/export?dataset=missing", headers=headers).status_code == 404


def test_excel_preserves_long_results_and_never_executes_source_text(cases):
    library = CaseLibraryService(cases)
    document = library.export(case_id=cases.case_ids()[0])
    document["cases"][0]["initial_information"]["demographics"]["age"] = [[52.0, "year"]]
    record = document["cases"][0]["actions"][0]["records"][0]
    record["result"] = {"text": "=HYPERLINK(\"bad\")\n" + "测试🩺" * 18000, "empty": None}
    workbook = load_workbook(BytesIO(excel_bytes(document)))
    assert len(workbook.sheetnames) == 9
    assert workbook["病例总览"].max_row == 2
    assert json.loads(workbook["病例总览"]["C2"].value) == [[52.0, "year"]]
    rows = [row for row in workbook["证据结果"].iter_rows(min_row=2, values_only=True)
            if row[1] == document["cases"][0]["actions"][0]["id"] and row[2] == 1]
    restored = json.loads("".join(row[-1] for row in sorted(rows, key=lambda r: r[6])))
    assert restored == record
    assert len(rows) > 1
    assert all(cell.data_type != "f" for ws in workbook for row in ws for cell in row)
    assert all(ws.freeze_panes == "A2" and ws.auto_filter.ref for ws in workbook)
