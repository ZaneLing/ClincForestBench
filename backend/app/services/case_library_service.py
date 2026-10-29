"""Read-only case library and lossless local JSON/Excel exports."""
from __future__ import annotations

from collections import Counter
from io import BytesIO
import json
from pathlib import Path

from backend.app.domain.evidence import human_readable_response
from backend.app.services.encounter_catalog import EncounterCatalog
from backend.app.services.localization_service import localize_payload, normalize_locale


class CaseLibraryService:
    def __init__(self, cases, catalog=None):
        self.cases = cases
        self.catalog = catalog or EncounterCatalog(cases)

    def index(self):
        rows = []
        for cid, visit in self.catalog.visits.items():
            case = self.cases.get(cid)
            counts = Counter(self.catalog.items[key]["kind"] for key in self.catalog.bindings[cid])
            rows.append({
                "case_id": cid, "dataset": visit["dataset"],
                "demographics": case.demographics.model_dump(mode="json"),
                "inquiry_count": counts["inquiry"], "operation_count": counts["operation"],
                "model_evidence_count": len(self.cases.catalog_for_case(cid)),
                "reference_quality": case.reference_quality,
            })
        return {"case_count": len(rows), "cases": rows,
                "items": list(self.catalog.items.values()),
                "datasets": dict(sorted(Counter(r["dataset"] for r in rows).items()))}

    def detail(self, case_id, locale="en"):
        locale = normalize_locale(locale)
        case = self.cases.get(case_id)
        inspection = self.catalog.inspection(case_id, locale)
        actions = []
        for item_id, records in inspection["bindings"].items():
            item = self.catalog.items[item_id]
            actions.append({
                **item, "display_name": item["label"][locale],
                "availability": "RECORDED", "record_count": len(records),
                "partial_panel": any(r.get("partial_panel", False) for r in records),
                "records": records,
            })
        evidence = []
        truth = case.truth_map()
        for eid, definition in self.cases.catalog_for_case(case_id).items():
            response = truth[eid]
            initial = eid == case.initial_evidence.evidence_id
            answer = response.values if response.values else response.value
            evidence.append({
                "evidence_id": eid,
                "question": localize_payload(definition.question_en, locale),
                "question_en": definition.question_en,
                "action_kind": definition.action_kind,
                "clinical_domain": definition.clinical_domain,
                "parent_evidence_id": definition.parent_evidence_id,
                "access": "INITIAL" if initial else "CONDITIONAL" if definition.parent_evidence_id else "DIRECT",
                "status": response.status.value,
                "data_type": definition.data_type.value,
                "answer": localize_payload(answer, locale),
                "display_answer": localize_payload(human_readable_response(response, definition), locale),
                "source_response": response.model_dump(mode="json"),
            })
        initial = inspection["visit"]
        media = {}

        def collect_media(value):
            if isinstance(value, dict):
                asset = value.get("media_asset")
                if isinstance(asset, str):
                    from urllib.parse import quote
                    media[asset] = {"asset": asset, "url": f"/cases/{quote(case_id, safe='')}/media/{quote(asset, safe='')}"}
                for child in value.values():
                    collect_media(child)
            elif isinstance(value, list):
                for child in value:
                    collect_media(child)

        collect_media(initial)
        collect_media(actions)
        return {
            "schema_version": "clincforestbench.case-library.v1",
            "locale": locale, "protocol": "UNORDERED_EVIDENCE",
            "case_id": case_id, "dataset": case.dataset.name,
            "case_hash": case.case_hash,
            "identity": {"dataset": case.dataset.model_dump(mode="json"),
                         "case_type": case.case_type, "generation_mode": case.generation_mode},
            "initial_information": initial,
            "actions": actions,
            "counts": inspection["counts"],
            "model_evidence": evidence,
            "treatments": inspection["treatments"],
            "media": list(media.values()),
            "source": case.source_provenance.get("source", {}),
            "mapping_audit": inspection["audit"],
            "reference": {"oracle": localize_payload(case.oracle.model_dump(mode="json"), locale),
                          "quality": case.reference_quality,
                          "fidelity_note": localize_payload(case.fidelity_note, locale)},
            "notes": {
                "initial_information": "Visit initial state; hidden reference data is not part of the initial presentation.",
                "actions": "Recorded Visit actions only. Other global catalog items have no recorded result for this case.",
                "model_evidence": "Canonical model action pool; INITIAL is already known and CONDITIONAL describes a parent evidence rule, not source chronology.",
                "results": "Fixed source-backed values. DEFAULT and NOT_APPLICABLE are distinct from negative findings.",
                "reference": "Research-only reference; provisional labels are not independently adjudicated diagnoses.",
            },
        }

    def export(self, locale="en", dataset=None, case_id=None):
        if case_id:
            ids = [case_id]
            self.cases.get(case_id)
        else:
            ids = [cid for cid, case in self.cases.cases.items()
                   if not dataset or case.dataset.name == dataset]
            if not ids:
                raise KeyError(f"Unknown dataset: {dataset}")
        inventory_path = self.cases.root / "data/manifests/dataset_inventory.json"
        inventory = json.loads(inventory_path.read_text()) if inventory_path.exists() else {}
        return {
            "schema_version": "clincforestbench.case-library-export.v1",
            "locale": normalize_locale(locale), "protocol": "UNORDERED_EVIDENCE",
            "scope": {"dataset": dataset, "case_id": case_id},
            "case_count": len(ids), "inventory": inventory,
            "action_catalog": list(self.catalog.items.values()),
            "cases": [self.detail(cid, locale) for cid in ids],
        }


def json_bytes(document):
    return (json.dumps(document, ensure_ascii=False, indent=2) + "\n").encode("utf-8")


def excel_bytes(document):
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.active)
    headers = {
        "使用说明": ["字段 / Field", "说明 / Description"],
        "数据集清单": ["数据集 / Dataset", "实际例数 / Actual", "目标例数 / Target", "缺口原因 / Shortfall"],
        "病例总览": ["Case ID", "数据集", "年龄", "性别", "病例类型", "生成方式", "Case Hash", "质量标记 JSON"],
        "初始信息": ["Case ID", "字段路径", "片段编号", "片段总数", "值类型", "完整值 JSON（按编号拼接）"],
        "动作目录": ["动作 ID", "类型", "分类", "英文名称", "中文名称"],
        "可选动作": ["Case ID", "数据集", "动作 ID", "类型", "分类", "名称", "记录数", "仅有部分分项"],
        "证据结果": ["Case ID", "动作 ID", "记录编号", "源记录 ID", "源名称", "映射方式", "片段编号", "片段总数", "完整记录 JSON（按编号拼接）"],
        "模型证据池": ["Case ID", "证据 ID", "问题", "动作类型", "领域", "父证据 ID", "访问规则", "结果状态", "数据类型", "片段编号", "片段总数", "完整回答 JSON（按编号拼接）"],
        "诊断参考": ["Case ID", "字段路径", "片段编号", "片段总数", "值类型", "完整值 JSON（按编号拼接）"],
    }
    sheets = {}
    for name, fields in headers.items():
        ws = workbook.create_sheet(name)
        ws.append(fields)
        sheets[name] = ws

    def parts(value):
        text = json.dumps(value, ensure_ascii=False)
        # Excel counts UTF-16 units; a smaller limit also handles supplementary
        # Unicode characters without losing content.
        return [text[i:i + 14000] for i in range(0, len(text), 14000)] or [""]

    def flattened(value, prefix=""):
        if isinstance(value, dict) and value:
            for key, child in value.items():
                yield from flattened(child, f"{prefix}.{key}" if prefix else key)
        else:
            yield prefix, value

    def append_values(sheet, cid, value):
        for path, item in flattened(value):
            chunks = parts(item)
            for number, chunk in enumerate(chunks, 1):
                sheet.append([cid, path, number, len(chunks), type(item).__name__, chunk])

    instructions = {
        "阅读顺序": "病例总览 → 初始信息 → 可选动作 → 证据结果。模型证据池使用独立的 canonical evidence ID。",
        "初始信息": "来自医生 Visit 初始状态；参考诊断不混入初始主诉。",
        "可选动作": "只列出本病例有固定结果的动作；动作目录是共享全集，未绑定的动作没有源结果。",
        "证据结果": "每条源记录原样导出为 JSON；同一记录的片段按片段编号拼接后可解析。",
        "模型证据池": "INITIAL=已知初始信息；DIRECT=直接询问；CONDITIONAL=父证据条件，不代表时间解锁。",
        "结果状态": "ABSENT=否定；DEFAULT=数据集默认；NOT_APPLICABLE=不适用；VALUE=固定结果；PRESENT=存在。",
        "诊断参考": "研究审核用途。质量标记需一并阅读，暂定标签不等于临床独立裁定。",
        "来源": "当前本地病例清单及确定性的问询/检查绑定，协议为 UNORDERED_EVIDENCE。",
        "文件范围": json.dumps(document["scope"], ensure_ascii=False),
        "长文本": "JSON 分片完整保留结果。表格仅限制可见行高；可在编辑栏查看或拼接全部正文。",
    }
    for row in instructions.items():
        sheets["使用说明"].append(row)
    inventory = document["inventory"]
    for dataset, actual in inventory.get("dataset_case_counts", {}).items():
        sheets["数据集清单"].append([dataset, actual, inventory.get("requested_cases_per_dataset"),
                                      inventory.get("shortfalls", {}).get(dataset, {}).get("reason", "")])
    for item in document["action_catalog"]:
        sheets["动作目录"].append([item["id"], item["kind"], item.get("category"),
                                   item["label"]["en"], item["label"]["zh-CN"]])
    for case in document["cases"]:
        cid = case["case_id"]
        demo = case["initial_information"]["demographics"]
        age = demo.get("age")
        if isinstance(age, (dict, list)):
            age = json.dumps(age, ensure_ascii=False)
        sheets["病例总览"].append([cid, case["dataset"], age, demo.get("sex"),
                                   case["identity"]["case_type"], case["identity"]["generation_mode"],
                                   case["case_hash"], json.dumps(case["reference"]["quality"], ensure_ascii=False)])
        append_values(sheets["初始信息"], cid, case["initial_information"])
        append_values(sheets["诊断参考"], cid, case["reference"])
        for action in case["actions"]:
            aid = action["id"]
            sheets["可选动作"].append([cid, case["dataset"], aid, action["kind"], action.get("category"),
                                       action["display_name"], action["record_count"], action["partial_panel"]])
            for index, record in enumerate(action["records"], 1):
                chunks = parts(record)
                for number, chunk in enumerate(chunks, 1):
                    sheets["证据结果"].append([cid, aid, index, record["source_id"], record.get("source_label"),
                                               record.get("mapping_method"), number, len(chunks), chunk])
        for evidence in case["model_evidence"]:
            chunks = parts({"answer": evidence["answer"], "display_answer": evidence["display_answer"],
                            "source_response": evidence["source_response"]})
            for number, chunk in enumerate(chunks, 1):
                sheets["模型证据池"].append([cid, evidence["evidence_id"], evidence["question"],
                                             evidence["action_kind"], evidence["clinical_domain"],
                                             evidence["parent_evidence_id"], evidence["access"], evidence["status"],
                                             evidence["data_type"], number, len(chunks), chunk])
    bottom = Border(bottom=Side(style="thin", color="D9DEE7"))
    for ws in workbook:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
            cell.fill = PatternFill("solid", fgColor="233B55")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 32
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value = ILLEGAL_CHARACTERS_RE.sub("", cell.value)
                    cell.data_type = "s"  # Source prose must never become an Excel formula.
                cell.font = Font(name="Arial", size=10, color="233044")
                cell.fill = PatternFill("solid", fgColor="F7F9FC" if cell.row % 2 == 0 else "FFFFFF")
                cell.border = bottom
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            ws.row_dimensions[row[0].row].height = 44
        for index, field in enumerate(headers[ws.title], 1):
            width = 85 if "JSON" in field or "说明" in field else 52 if "名称" in field or "问题" in field else 29
            ws.column_dimensions[get_column_letter(index)].width = width
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
